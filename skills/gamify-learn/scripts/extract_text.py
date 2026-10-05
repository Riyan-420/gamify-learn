#!/usr/bin/env python3
"""Pull text (and, for .pptx, images) out of course material so a course.json can be written from it.

usage: python extract_text.py <file-or-folder> [-o out_dir] [--images]

Supports: .pdf .pptx .docx .txt .md .html
Optional packages (install only what you need):  pip install pypdf python-pptx python-docx
Writes one <name>.txt per input to out_dir (default ./extracted). With --images, .pptx pictures are
saved to out_dir/<name>_images/ named by slide number, so you can reuse real figures in 'figure' steps
(only if you have the right to reuse them).
"""
import argparse, os, re, sys, zipfile


def need(pkg, pipname=None):
    sys.exit(f"missing package '{pkg}'. Install it:  pip install {pipname or pkg}")


def pdf_text(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        need("pypdf")
    r = PdfReader(path)
    return "\n\n".join(f"--- page {i} ---\n{(p.extract_text() or '').strip()}" for i, p in enumerate(r.pages, 1))


def pptx_text(path, img_dir=None):
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        need("python-pptx")
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        k = 0
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            if img_dir and sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                os.makedirs(img_dir, exist_ok=True)
                k += 1
                ext = sh.image.ext or "png"
                with open(os.path.join(img_dir, f"s{i:02d}_{k}.{ext}"), "wb") as f:
                    f.write(sh.image.blob)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        out.append(f"--- slide {i} ---\n" + "\n".join(parts) + (f"\n[notes] {notes}" if notes else ""))
    return "\n\n".join(out)


def docx_text(path):
    try:
        import docx
    except ImportError:
        need("docx", "python-docx")
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs if p.text.strip())


def html_text(path):
    s = open(path, encoding="utf-8", errors="replace").read()
    s = re.sub(r"(?is)<(script|style).*?</\1>", "", s)
    s = re.sub(r"(?s)<[^>]+>", " ", s)
    return re.sub(r"[ \t]+", " ", re.sub(r"\n\s*\n+", "\n\n", s))


def extract(path, out_dir, images):
    ext = os.path.splitext(path)[1].lower()
    name = os.path.splitext(os.path.basename(path))[0]
    if ext == ".pdf":
        t = pdf_text(path)
    elif ext == ".pptx":
        t = pptx_text(path, os.path.join(out_dir, name + "_images") if images else None)
    elif ext == ".docx":
        t = docx_text(path)
    elif ext in (".txt", ".md"):
        t = open(path, encoding="utf-8", errors="replace").read()
    elif ext in (".html", ".htm"):
        t = html_text(path)
    else:
        return None
    os.makedirs(out_dir, exist_ok=True)
    dest = os.path.join(out_dir, name + ".txt")
    with open(dest, "w", encoding="utf-8") as f:
        f.write(t)
    print(f"{path} -> {dest}  ({len(t):,} chars)")
    if len(t.strip()) < 200 and ext == ".pdf":
        print("  note: almost no text found; this PDF is probably scanned images (needs OCR).")
    return dest


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("-o", "--out", default="extracted")
    ap.add_argument("--images", action="store_true")
    a = ap.parse_args()
    files = []
    if os.path.isdir(a.src):
        for r, _, fs in os.walk(a.src):
            files += [os.path.join(r, f) for f in sorted(fs)]
    else:
        files = [a.src]
    n = sum(1 for f in files if extract(f, a.out, a.images))
    if not n:
        sys.exit("nothing extracted (supported: .pdf .pptx .docx .txt .md .html)")


if __name__ == "__main__":
    main()
